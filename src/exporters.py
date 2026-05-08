import io
import re


def markdown_to_pdf_bytes(markdown_text: str, title: str = "Report") -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "", 11)

    def _line(text: str, size: int = 11, bold: bool = False, indent: str = "") -> None:
        pdf.set_font("Helvetica", "B" if bold else "", size)
        pdf.multi_cell(
            0, max(6, int(size * 0.6)),
            indent + _safe(text),
            new_x=XPos.LMARGIN, new_y=YPos.NEXT,
        )

    for line in markdown_text.splitlines():
        stripped = line.rstrip()
        if not stripped:
            pdf.ln(3)
            continue
        m_h = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if m_h:
            level = len(m_h.group(1))
            _line(m_h.group(2), size=max(11, 16 - level), bold=True)
            continue
        if stripped.startswith(("- ", "* ", "+ ")):
            _line(stripped[2:], indent="  - ")
            continue
        m_num = re.match(r"^\d+\.\s+(.+)$", stripped)
        if m_num:
            _line(stripped, indent="  ")
            continue
        _line(stripped)

    out = pdf.output(dest="S")
    return bytes(out) if isinstance(out, (bytes, bytearray)) else out.encode("latin-1")


def markdown_to_docx_bytes(markdown_text: str, title: str = "Report") -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Helvetica"
    style.font.size = Pt(11)

    for line in markdown_text.splitlines():
        stripped = line.rstrip()
        if not stripped:
            doc.add_paragraph("")
            continue
        m_h = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if m_h:
            level = len(m_h.group(1))
            doc.add_heading(m_h.group(2), level=min(level, 4))
            continue
        if stripped.startswith(("- ", "* ", "+ ")):
            p = doc.add_paragraph(stripped[2:], style="List Bullet")
            continue
        m_num = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if m_num:
            doc.add_paragraph(m_num.group(2), style="List Number")
            continue
        doc.add_paragraph(stripped)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _safe(text: str) -> str:
    return text.encode("latin-1", "replace").decode("latin-1")
